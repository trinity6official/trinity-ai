"""Document extraction for Trinity personal knowledge.

Text/code files retain line references. PDFs retain page references. DOCX files
retain paragraph references. Optional document libraries are imported lazily so
Android text/code indexing keeps working even when rich-document packages are
not installed.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Iterable


TEXT_SUFFIXES = {
    ".txt", ".md", ".rst", ".py", ".json", ".yaml", ".yml", ".toml",
    ".ini", ".cfg", ".csv", ".tsv", ".sh", ".ps1", ".rb", ".go",
    ".js", ".ts", ".tsx", ".jsx", ".html", ".htm", ".css", ".sql",
    ".xml", ".java", ".c", ".h", ".cpp", ".hpp", ".rs", ".swift",
    ".kt", ".kts",
}
TEXT_NAMES = {
    "readme", "license", "makefile", "dockerfile", "gemfile", "rakefile",
    "procfile",
}
RICH_SUFFIXES = {".pdf", ".docx"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | RICH_SUFFIXES
SUPPORTED_NAMES = TEXT_NAMES


@dataclass(frozen=True)
class ExtractedChunk:
    content: str
    locator_kind: str
    start_locator: int
    end_locator: int

    def source_ref(self, path: str | Path) -> str:
        source = str(path)
        if self.locator_kind == "page":
            if self.start_locator == self.end_locator:
                return f"{source}#p{self.start_locator}"
            return f"{source}#p{self.start_locator}-p{self.end_locator}"
        if self.locator_kind == "paragraph":
            if self.start_locator == self.end_locator:
                return f"{source}#P{self.start_locator}"
            return f"{source}#P{self.start_locator}-P{self.end_locator}"
        return f"{source}#L{self.start_locator}-L{self.end_locator}"


class DocumentExtractor:
    def __init__(self, *, chunk_chars: int = 1600, overlap_lines: int = 2) -> None:
        self.chunk_chars = max(400, int(chunk_chars))
        self.overlap_lines = max(0, int(overlap_lines))

    @staticmethod
    def supported(path: Path) -> bool:
        return path.suffix.lower() in SUPPORTED_SUFFIXES or path.name.lower() in SUPPORTED_NAMES

    @staticmethod
    def document_type(path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return "pdf"
        if suffix == ".docx":
            return "docx"
        return "text"

    @staticmethod
    def _decode_text(raw: bytes) -> str:
        if b"\x00" in raw[:8192]:
            raise ValueError("binary file")
        try:
            return raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            return raw.decode("utf-8", errors="replace")

    def extract_bytes(self, path: Path, raw: bytes) -> list[ExtractedChunk]:
        if path.suffix.lower() in RICH_SUFFIXES:
            raise ValueError("rich documents must be extracted from their file path")
        return self._line_chunks(self._decode_text(raw))

    def extract(self, path: Path) -> list[ExtractedChunk]:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._extract_pdf(path)
        if suffix == ".docx":
            return self._extract_docx(path)
        return self.extract_bytes(path, path.read_bytes())

    def _line_chunks(self, text: str) -> list[ExtractedChunk]:
        lines = text.splitlines()
        if not lines:
            return []
        chunks: list[ExtractedChunk] = []
        start = 0
        total = len(lines)
        while start < total:
            end = start
            chars = 0
            while end < total:
                next_size = len(lines[end]) + 1
                if end > start and chars + next_size > self.chunk_chars:
                    break
                chars += next_size
                end += 1
            if end == start:
                end += 1
            content = "\n".join(lines[start:end]).strip()
            if content:
                chunks.append(ExtractedChunk(content, "line", start + 1, end))
            if end >= total:
                break
            next_start = end - self.overlap_lines
            if next_start <= start:
                next_start = start + 1
            start = next_start
        return chunks

    def _split_block(self, text: str, locator_kind: str, locator: int) -> list[ExtractedChunk]:
        clean = text.strip()
        if not clean:
            return []
        if len(clean) <= self.chunk_chars:
            return [ExtractedChunk(clean, locator_kind, locator, locator)]

        lines = clean.splitlines() or [clean]
        chunks: list[ExtractedChunk] = []
        current: list[str] = []
        size = 0
        for line in lines:
            if current and size + len(line) + 1 > self.chunk_chars:
                chunks.append(
                    ExtractedChunk("\n".join(current).strip(), locator_kind, locator, locator)
                )
                current = []
                size = 0
            if len(line) > self.chunk_chars and not current:
                for offset in range(0, len(line), self.chunk_chars):
                    part = line[offset : offset + self.chunk_chars].strip()
                    if part:
                        chunks.append(ExtractedChunk(part, locator_kind, locator, locator))
                continue
            current.append(line)
            size += len(line) + 1
        if current:
            chunks.append(ExtractedChunk("\n".join(current).strip(), locator_kind, locator, locator))
        return [chunk for chunk in chunks if chunk.content]

    def _extract_pdf(self, path: Path) -> list[ExtractedChunk]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - environment dependent
            raise RuntimeError("PDF support requires pypdf (install requirements.txt)") from exc

        reader = PdfReader(str(path))
        chunks: list[ExtractedChunk] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            chunks.extend(self._split_block(text, "page", page_number))
        return chunks

    def _extract_docx(self, path: Path) -> list[ExtractedChunk]:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - environment dependent
            raise RuntimeError("DOCX support requires python-docx (install requirements.txt)") from exc

        document = Document(str(path))
        chunks: list[ExtractedChunk] = []
        for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
            text = (paragraph.text or "").strip()
            if text:
                chunks.extend(self._split_block(text, "paragraph", paragraph_number))

        # Tables are appended with stable synthetic paragraph locations after the
        # ordinary paragraphs. This keeps citations deterministic without claiming
        # page numbers DOCX does not reliably expose.
        locator = len(document.paragraphs) + 1
        for table in document.tables:
            for row in table.rows:
                text = " | ".join((cell.text or "").strip() for cell in row.cells).strip(" |")
                if text:
                    chunks.extend(self._split_block(text, "paragraph", locator))
                    locator += 1
        return chunks

    def read_reference(self, source_ref: str) -> dict[str, object]:
        if "#" not in source_ref:
            raise ValueError("Source reference must include #L, #p, or #P locator")
        raw_path, fragment = source_ref.rsplit("#", 1)
        path = Path(raw_path).expanduser().resolve()
        match = re.fullmatch(r"([LpP])(\d+)(?:-([LpP])?(\d+))?", fragment)
        if not match:
            raise ValueError(f"Invalid knowledge source reference: {source_ref}")

        prefix = match.group(1)
        start = int(match.group(2))
        end = int(match.group(4) or start)
        if end < start:
            start, end = end, start

        if prefix == "L":
            raw = path.read_bytes()
            text = self._decode_text(raw)
            lines = text.splitlines()
            end = min(end, len(lines))
            content = "\n".join(lines[max(0, start - 1):end])
            return {
                "source_path": str(path),
                "source_ref": f"{path}#L{start}-L{end}",
                "locator_kind": "line",
                "start_locator": start,
                "end_locator": end,
                "content": content,
            }

        if prefix == "p":
            try:
                from pypdf import PdfReader
            except ImportError as exc:  # pragma: no cover
                raise RuntimeError("PDF support requires pypdf") from exc
            reader = PdfReader(str(path))
            end = min(end, len(reader.pages))
            content = "\n\n".join(
                (reader.pages[index - 1].extract_text() or "").strip()
                for index in range(start, end + 1)
                if 1 <= index <= len(reader.pages)
            )
            ref = f"{path}#p{start}" if start == end else f"{path}#p{start}-p{end}"
            return {
                "source_path": str(path), "source_ref": ref,
                "locator_kind": "page", "start_locator": start,
                "end_locator": end, "content": content,
            }

        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("DOCX support requires python-docx") from exc
        document = Document(str(path))
        paragraphs = [(p.text or "") for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                text = " | ".join((cell.text or "").strip() for cell in row.cells).strip(" |")
                if text:
                    paragraphs.append(text)
        end = min(end, len(paragraphs))
        content = "\n".join(paragraphs[max(0, start - 1):end])
        ref = f"{path}#P{start}" if start == end else f"{path}#P{start}-P{end}"
        return {
            "source_path": str(path), "source_ref": ref,
            "locator_kind": "paragraph", "start_locator": start,
            "end_locator": end, "content": content,
        }
